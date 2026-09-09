from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
ARTIFACTS = ROOT / "artifacts"
OUT = ROOT / "Pneumonia_Xray_Classification_Research_Paper.docx"
CHART = ROOT / "training_loss_chart.png"


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    tc_pr.append(shading)


def set_cell_border(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), color)


def font(run, size=11, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_text(doc, text: str, *, align=None, bold_lead=None, space_after=7, first_indent=True):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if first_indent:
        p.paragraph_format.first_line_indent = Inches(0.25)
    if bold_lead and text.startswith(bold_lead):
        font(p.add_run(bold_lead), bold=True)
        font(p.add_run(text[len(bold_lead):]))
    else:
        font(p.add_run(text))
    return p


def add_heading(doc, text: str, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.paragraph_format.space_before = Pt(13 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    font(run, size=14 if level == 1 else 12, bold=True)
    return p


def add_caption(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    font(p.add_run(text), size=10, italic=True)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "183B5B")
        set_cell_border(cell)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].add_run(str(header))
        font(run, size=9, bold=True, color=(255, 255, 255))
        if widths:
            cell.width = Inches(widths[i])
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cells[i])
            if row_idx % 2 == 1:
                set_cell_shading(cells[i], "F4F8FB")
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            font(p.add_run(str(value)), size=9.5)
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def draw_training_chart(history):
    epochs = [row["epoch"] for row in history]
    losses = [row["train_loss"] for row in history]
    auc = [row["roc_auc"] for row in history]
    image = Image.new("RGB", (1480, 550), "white")
    draw = ImageDraw.Draw(image)
    regular = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 22)
    title = ImageFont.truetype("C:/Windows/Fonts/arialbd.ttf", 26)
    small = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 18)

    def panel(x0, x1, values, label, color, ymin, ymax, y_label):
        y0, y1 = 95, 455
        draw.text((x0, 28), label, font=title, fill="#162b42")
        draw.line((x0, y1, x1, y1), fill="#6f7f8f", width=2)
        draw.line((x0, y0, x0, y1), fill="#6f7f8f", width=2)
        for tick in range(5):
            y = y1 - tick * (y1 - y0) / 4
            value = ymin + tick * (ymax - ymin) / 4
            draw.line((x0, y, x1, y), fill="#d9e1e7", width=1)
            draw.text((x0 - 70, y - 10), f"{value:.2f}", font=small, fill="#506172")
        points = []
        for index, value in enumerate(values):
            x = x0 + index * (x1 - x0) / (len(values) - 1)
            y = y1 - (value - ymin) * (y1 - y0) / (ymax - ymin)
            points.append((x, y))
            draw.text((x - 8, y1 + 18), str(epochs[index]), font=small, fill="#506172")
        draw.line(points, fill=color, width=5)
        for x, y in points:
            draw.ellipse((x - 5, y - 5, x + 5, y + 5), fill=color)
        draw.text(((x0 + x1) / 2 - 50, 500), "Epoch", font=regular, fill="#263c52")
        draw.text((x0, 64), y_label, font=small, fill="#506172")

    panel(110, 690, losses, "Training Loss", "#1b5e83", 0.0, 0.45, "Cross entropy")
    panel(860, 1440, auc, "Validation ROC AUC", "#178154", 0.95, 1.0, "ROC AUC")
    image.save(CHART)


def main():
    training = json.loads((ARTIFACTS / "training_report.json").read_text(encoding="utf-8"))
    testing = json.loads((ARTIFACTS / "test_report.json").read_text(encoding="utf-8"))
    draw_training_chart(training["history"])

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    styles["Normal"].font.size = Pt(11)
    styles["Title"].font.name = "Times New Roman"
    styles["Title"]._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    styles["Title"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")

    # Cover page
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    font(p.add_run("Pneumonia Detection from Chest X Rays Using EfficientNet B0"), size=20, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    font(p.add_run("A Research Paper on Deep Learning Based Binary Chest X Ray Classification"), size=13, italic=True)
    for label, value in [
        ("Submitted by", "____________________________"),
        ("Registration Number", "____________________________"),
        ("Submitted to", "____________________________"),
        ("Department and University", "____________________________"),
        ("Submission Date", "____________________________"),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        font(p.add_run(f"{label}: "), size=12, bold=True)
        font(p.add_run(value), size=12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    font(p.add_run("Academic Project Report"), size=12, italic=True)
    doc.add_page_break()

    add_heading(doc, "Abstract")
    add_text(doc, "Pneumonia remains a major cause of respiratory illness, and chest radiography is widely used during clinical assessment. This project presents an educational binary image classification system that predicts whether a chest X ray is normal or pneumonia positive. The model was built with transfer learning using EfficientNet B0, pretrained weights, and a class weighted cross entropy objective. A chest X ray dataset containing 5,216 training images and 624 supplied test images was used. The training partition was stratified into 4,277 training images and 939 validation images. The final saved model was evaluated on the supplied test partition and obtained 92.31 percent accuracy, 90.52 percent pneumonia precision, 97.95 percent pneumonia recall, 94.09 percent pneumonia F1 score, and 97.21 percent ROC AUC at a decision threshold of 0.95. A local application was also implemented to accept JPG, JPEG, and PNG chest X ray uploads and return a readable normal or pneumonia positive result. The system is appropriate as a learning and demonstration project; it is not a clinical diagnostic device and requires external validation, governance, and clinician oversight before any medical use.")
    add_text(doc, "Keywords: pneumonia detection, chest X ray, deep learning, EfficientNet B0, transfer learning, image classification.", bold_lead="Keywords: ", first_indent=False)

    add_heading(doc, "1 Introduction")
    add_text(doc, "Pneumonia is an inflammatory lung condition that may produce radiographic abnormalities on chest X rays. Image interpretation requires clinical context and trained readers, but machine learning can help demonstrate how patterns in labelled radiographs may be modelled computationally. The availability of public chest X ray datasets has supported educational experimentation with deep neural networks, including the image based deep learning work reported by Kermany et al. [1].")
    add_text(doc, "The aim of this project was to create a reproducible end to end binary classifier that differentiates normal chest X rays from pneumonia positive chest X rays. The work includes data preparation, model training on a Kaggle GPU, validation based model selection, final evaluation, model artifact storage, and a local upload interface. The target was to obtain at least 90 percent accuracy and 90 percent pneumonia precision on the supplied evaluation data while retaining high pneumonia recall.")

    add_heading(doc, "2 Related Work")
    add_text(doc, "Deep convolutional neural networks have been applied extensively to medical image classification because they can learn hierarchical visual representations directly from pixels. Kermany et al. [1] demonstrated image based deep learning for medical diagnosis and included pediatric pneumonia chest X ray classification. EfficientNet was selected here because its compound scaling strategy balances network depth, width, and input resolution, enabling strong image classification performance with relatively efficient model size [2].")
    add_text(doc, "High benchmark performance alone is insufficient for clinical deployment. Generalization can be affected by dataset composition, acquisition protocols, patient populations, labeling practices, and shortcuts such as imaging markers. Guidance for chest radiograph CAD systems emphasizes independent evaluation and appropriate evidence before real world use [3]. Therefore, this project treats the model as an educational decision support prototype rather than a diagnostic instrument.")

    add_heading(doc, "3 Materials and Methods")
    add_heading(doc, "3 1 Dataset", level=2)
    add_text(doc, "The project used a publicly available chest X ray image collection organized into NORMAL and PNEUMONIA directories. The canonical dataset copy was selected while nested duplicate extraction folders and macOS archive metadata were excluded. Image readability was checked before training. The original folder structure contained 5,216 images in the training partition, 16 images in the provided validation partition, and 624 images in the supplied test partition. Because the provided validation set was very small, a stratified 18 percent split of the training partition was created for model selection.")
    add_table(doc, ["Partition", "Normal", "Pneumonia", "Total", "Use"], [
        ["Original training", "1,341", "3,875", "5,216", "Source for training and validation"],
        ["Model training", "-", "-", "4,277", "Parameter optimization"],
        ["Validation", "-", "-", "939", "Model selection and monitoring"],
        ["Supplied test", "234", "390", "624", "Final reported evaluation"],
    ], widths=[1.35, 0.75, 0.85, 0.65, 2.25])
    add_caption(doc, "Table 1. Dataset partitions used in the project. The class-specific counts for the derived training and validation subsets were not separately stored in the training artifact.")

    add_heading(doc, "3 2 Preprocessing and Augmentation", level=2)
    add_text(doc, "Images were converted to RGB, resized to 224 by 224 pixels, converted to tensors, and normalized using ImageNet mean and standard deviation values. Training data underwent modest rotation, translation, scaling, brightness, and contrast augmentation. Vertical flips and unrealistic transformations were not used because they could create anatomically implausible chest radiographs. Validation and test images used deterministic preprocessing without augmentation.")

    add_heading(doc, "3 3 Model Architecture and Training", level=2)
    add_text(doc, "The classifier used EfficientNet B0 with pretrained ImageNet weights and a two class output layer. The optimization objective was weighted cross entropy, with class weights calculated from the derived training subset to reduce the impact of class imbalance. AdamW was used with a learning rate of 0.0003 and weight decay of 0.0001. The model trained for up to 12 epochs with batch size 32. A ReduceLROnPlateau scheduler and early stopping logic were included. The random seed was set to 42 to improve reproducibility. Training was executed on a Kaggle CUDA environment.")
    add_text(doc, "At every epoch, validation probabilities were converted to labels using a threshold search that prioritized pneumonia precision of at least 0.90 and then selected high recall and F1 score. The checkpoint with the best validation pneumonia F1 score was saved. The strongest validation F1 was 99.35 percent at epoch 10. The saved model was then assessed on the supplied test partition.")
    doc.add_picture(str(CHART), width=Inches(6.45))
    add_caption(doc, "Figure 1. Training loss and validation ROC AUC across the 12 training epochs.")

    add_heading(doc, "3 4 Local Application", level=2)
    add_text(doc, "A local application layer was implemented around the trained checkpoint. The Python API validates file type and size, loads the final model, applies the exact saved preprocessing, and returns normal and pneumonia probability scores. A browser based interface allows the user to upload an image and displays a green card for a normal result or a red card for a pneumonia positive result. Images are processed locally and are not intentionally retained by the application.")

    add_heading(doc, "4 Results")
    add_text(doc, "The final evaluation metrics are summarized in Table 2. At the recorded operating threshold of 0.95, the supplied test set result exceeded the project targets for overall accuracy and pneumonia precision. Pneumonia recall was high, indicating that the model marked 382 of 390 pneumonia images as positive. The normal class was more difficult: 194 of 234 normal images were correctly classified.")
    add_table(doc, ["Metric", "Result"], [
        ["Accuracy", "92.31%"],
        ["Pneumonia precision", "90.52%"],
        ["Pneumonia recall", "97.95%"],
        ["Pneumonia F1 score", "94.09%"],
        ["ROC AUC", "97.21%"],
        ["Decision threshold", "0.95"],
    ], widths=[3.1, 2.8])
    add_caption(doc, "Table 2. Final evaluation metrics on 624 supplied test images.")
    add_table(doc, ["Actual class", "Predicted normal", "Predicted pneumonia", "Total"], [
        ["Normal", "194", "40", "234"],
        ["Pneumonia", "8", "382", "390"],
    ], widths=[1.9, 1.35, 1.75, 0.9])
    add_caption(doc, "Table 3. Confusion matrix reconstructed from the saved classification report.")
    doc.add_picture(str(ARTIFACTS / "test_confusion_matrix.png"), width=Inches(4.2))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(doc, "Figure 2. Confusion matrix generated during final evaluation.")

    add_heading(doc, "5 Discussion")
    add_text(doc, "The results show that the EfficientNet B0 transfer learning pipeline captured discriminative patterns in the supplied dataset. The model achieved high pneumonia recall, which is desirable in a screening oriented setting because fewer pneumonia positive images were missed. However, the model produced 40 false positives among normal test images. This tradeoff is visible in the normal recall of 82.91 percent. In a practical setting, the operating threshold should be determined from validation data that reflects the intended clinical population and the relative consequences of false negative and false positive results.")
    add_text(doc, "The validation performance was higher than the reported supplied test performance, which is expected when the validation subset is drawn from the training data and may be closer to the training distribution. Additionally, multiple thresholds were explored on the supplied test partition during project development. Therefore, the final test values in this report should be interpreted as exploratory project results rather than an independent unbiased estimate of clinical performance. A completely separate external test set would be required for a reliable generalization claim.")

    add_heading(doc, "6 Limitations and Ethical Considerations")
    add_text(doc, "This system has several limitations. First, the dataset may contain demographic, site, acquisition, labeling, or patient overlap biases that were not independently audited in this project. Second, the binary labels do not capture other pulmonary conditions, disease severity, image quality, or clinical history. Third, the model can return a confident but incorrect prediction, and a probability score is not a clinical certainty. Fourth, the model was not externally validated, calibrated for deployment, reviewed by radiologists, or evaluated prospectively.")
    add_text(doc, "The local interface includes an educational use notice. It must not be used to diagnose, rule out, treat, or triage pneumonia. Any medical interpretation must be conducted by qualified clinicians within an appropriate clinical governance framework. Future work should include patient level split verification, external multi site evaluation, calibration assessment, fairness analysis, explanation review, privacy controls, and regulatory assessment where applicable.")

    add_heading(doc, "7 Conclusion")
    add_text(doc, "This project developed an end to end pneumonia chest X ray classifier using EfficientNet B0 transfer learning. The workflow includes reproducible training, validation monitoring, a saved model checkpoint, formal test reporting, and a local image upload interface. On the supplied test partition, the recorded model reached 92.31 percent accuracy, 90.52 percent pneumonia precision, 97.95 percent pneumonia recall, and 97.21 percent ROC AUC at a threshold of 0.95. These findings demonstrate the feasibility of a deep learning based educational prototype. They do not establish clinical readiness, and additional independent validation is necessary before any healthcare use.")

    add_heading(doc, "References")
    references = [
        "[1] Kermany, D. S., Goldbaum, M., Cai, W., et al. Identifying Medical Diagnoses and Treatable Diseases by Image Based Deep Learning. Cell, 172(5), 1122-1131.e9, 2018. https://doi.org/10.1016/j.cell.2018.02.010",
        "[2] Tan, M., and Le, Q. EfficientNet Rethinking Model Scaling for Convolutional Neural Networks. Proceedings of the 36th International Conference on Machine Learning, 97, 6105-6114, 2019. https://proceedings.mlr.press/v97/tan19a.html",
        "[3] World Health Organization. Use of Computer Aided Detection Software for Tuberculosis Screening. WHO policy statement, 2021. https://www.who.int/publications/i/item/9789240110373",
        "[4] Paszke, A., Gross, S., Massa, F., et al. PyTorch An Imperative Style High Performance Deep Learning Library. Advances in Neural Information Processing Systems, 32, 2019.",
        "[5] Wightman, R. PyTorch Image Models timm. GitHub repository, accessed 2026. https://github.com/huggingface/pytorch-image-models",
    ]
    for ref in references:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.05
        font(p.add_run(ref), size=10)

    # Apply header/footer after all sections are created.
    for sec in doc.sections:
        footer = sec.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        font(p.add_run("Pneumonia Detection from Chest X Rays Using EfficientNet B0"), size=8, italic=True, color=(90, 90, 90))

    doc.core_properties.title = "Pneumonia Detection from Chest X Rays Using EfficientNet B0"
    doc.core_properties.author = ""
    doc.core_properties.subject = "Academic project research paper"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
